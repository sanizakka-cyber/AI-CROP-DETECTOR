"""
Exports every farmer-facing string that is new or changed between two git
refs into one spreadsheet for native-speaker review (Hausa, Yoruba, Igbo,
French, Arabic). Reviewers fill the "OK / correction" column; import with
import_review_sheet.py.

  py -I scripts/i18n/export_review_sheet.py --base origin/main --head origin/scan-error-codes \
      --out docs/review/translation-review-2026-10.xlsx

Sources: web msas-system/lang/<lang>.json (7 languages), app
mobile/lib/i18n.js and mobile/lib/scanErrors.js (6 languages; no Arabic).
Staff-only labels (Technical details box) are left out.
"""
import argparse
import json
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
LANGS = [("ha", "Hausa"), ("yo", "Yoruba"), ("ig", "Igbo"), ("fr", "French"), ("ar", "Arabic")]
STAFF_ONLY = {"Technical details", "Copy details", "Copied", "Error code", "HTTP status", "Failed service", "Error message", "Time",
              "technical", "share"}


def git_show(ref, path):
    r = subprocess.run(["git", "show", f"{ref}:{path}"], capture_output=True)
    return r.stdout.decode("utf-8") if r.returncode == 0 else None


def web_strings(ref):
    out = {}
    for lang in ["en", "ha", "yo", "ig", "fr", "ar", "ff"]:
        text = git_show(ref, f"msas-system/lang/{lang}.json")
        out[lang] = json.loads(text) if text else {}
    return out


def app_strings(ref, path, kind):
    text = git_show(ref, path)
    if text is None:
        return {}
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
        f.write(text)
        tmp = f.name
    try:
        r = subprocess.run(["node", os.path.join(HERE, "dump_app_strings.mjs"), kind, tmp], capture_output=True)
        if r.returncode != 0:
            sys.exit(f"node failed on {path}: {r.stderr.decode()[:300]}")
        return json.loads(r.stdout.decode("utf-8"))
    finally:
        os.unlink(tmp)


def changed_rows(source, base, head, keys_are_english):
    """Rows for keys that are new in head or whose text changed in any language."""
    rows = []
    en_head = head.get("en", {})
    for key, en in en_head.items():
        if key in STAFF_ONLY or key.split(".")[-1] in STAFF_ONLY:
            continue
        is_new = key not in base.get("en", {})
        is_changed = any(head.get(l, {}).get(key) != base.get(l, {}).get(key) for l in head)
        if not (is_new or is_changed):
            continue
        rows.append({
            "source": source, "key": key, "english": en, "status": "new" if is_new else "changed",
            **{lang: head.get(lang, {}).get(key) for lang, _ in LANGS},
        })
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="origin/main")
    ap.add_argument("--head", default="origin/scan-error-codes")
    ap.add_argument("--out", default="docs/review/translation-review-2026-10.xlsx")
    a = ap.parse_args()

    rows = changed_rows("web", web_strings(a.base), web_strings(a.head), True)
    rows += changed_rows("app", app_strings(a.base, "mobile/lib/i18n.js", "i18n"), app_strings(a.head, "mobile/lib/i18n.js", "i18n"), False)
    rows += changed_rows("app-errors", app_strings(a.base, "mobile/lib/scanErrors.js", "scan"), app_strings(a.head, "mobile/lib/scanErrors.js", "scan"), False)

    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    for line in [
        ["MSAS FarmAI: translation review"],
        [f"Strings new or changed between {a.base} and {a.head} that farmers see (web and app)."],
        ["For each language: write OK if the text is right, or type the corrected text. Leave blank if you did not review it."],
        ["Keep :limit, :date, :field and :ref exactly as they are: the app fills them in."],
        ["Arabic is web only. Fulfulde is not in this sheet: it needs a translator first."],
        ["Do not change the Source, Key or English columns."],
    ]:
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Strings")
    header = ["Source", "Key", "Status", "English"]
    for _, name in LANGS:
        header += [f"{name} (current)", f"{name}: reviewer OK / correction"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        line = [r["source"], r["key"], r["status"], r["english"]]
        for lang, _ in LANGS:
            current = r[lang]
            if r["source"] != "web" and lang == "ar":
                line += ["(not in the app)", "n/a"]
            else:
                line += [current if current is not None else "(missing)", ""]
        ws.append(line)
    widths = [10, 34, 9, 48] + [42, 30] * len(LANGS)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    for i in range(len(LANGS)):
        col = ws.cell(row=1, column=6 + 2 * i).column_letter
        for row in range(2, ws.max_row + 1):
            ws[f"{col}{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"OK"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)

    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    wb.save(a.out)
    by_source = {}
    for r in rows:
        by_source[r["source"]] = by_source.get(r["source"], 0) + 1
    print(f"{len(rows)} strings -> {a.out}  {by_source}")


if __name__ == "__main__":
    main()
