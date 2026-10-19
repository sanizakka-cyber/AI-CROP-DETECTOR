"""
Applies reviewers' corrections from the translation review sheet
(export_review_sheet.py) back into the source files.

  py -I scripts/i18n/import_review_sheet.py docs/review/translation-review-2026-10-reviewed.xlsx [--dry-run]

A cell that is empty or "OK" changes nothing. Any other text replaces the
current translation:
  - web rows: msas-system/lang/<lang>.json, by key;
  - app rows: the old translated string in mobile/lib/i18n.js or
    mobile/lib/scanErrors.js, replaced only where it appears exactly once
    (anything else is listed for a manual fix, never guessed).
Placeholders (:limit, :date, :field, :ref) must survive; a correction that
drops one is refused. Prints every change.
"""
import argparse
import json
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
LANGS = [("ha", "Hausa"), ("yo", "Yoruba"), ("ig", "Igbo"), ("fr", "French"), ("ar", "Arabic")]
APP_FILES = {"app": "mobile/lib/i18n.js", "app-errors": "mobile/lib/scanErrors.js"}


def placeholders(text):
    return sorted(re.findall(r":[a-z_]+\b", text or ""))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sheet")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    ws = load_workbook(a.sheet)["Strings"]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    web = {}      # lang -> {key: new}
    app = []      # (file, lang, key, old, new)
    problems = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        source, key, english = row[col["Source"]], row[col["Key"]], row[col["English"]]
        for lang, name in LANGS:
            current = row[col[f"{name} (current)"]]
            review = row[col[f"{name}: reviewer OK / correction"]]
            if review is None or str(review).strip() == "" or str(review).strip().upper() in ("OK", "N/A"):
                continue
            new = str(review).strip()
            if placeholders(new) != placeholders(english):
                problems.append(f"{source} {key} [{lang}]: placeholders {placeholders(english)} missing or changed in the correction")
                continue
            if source == "web":
                web.setdefault(lang, {})[key] = new
            elif source in APP_FILES:
                app.append((APP_FILES[source], lang, key, current, new))

    for lang, changes in web.items():
        path = os.path.join(ROOT, "msas-system", "lang", f"{lang}.json")
        raw = open(path, encoding="utf-8", newline="").read()
        data = json.loads(raw)
        for key, new in changes.items():
            if key not in data:
                problems.append(f"web {key} [{lang}]: key not in {lang}.json")
                continue
            print(f"web [{lang}] {key!r}: {data[key]!r} -> {new!r}")
            old_line = f"{json.dumps(key, ensure_ascii=False)}: {json.dumps(data[key], ensure_ascii=False)}"
            new_line = f"{json.dumps(key, ensure_ascii=False)}: {json.dumps(new, ensure_ascii=False)}"
            if raw.count(old_line) != 1:
                problems.append(f"web {key} [{lang}]: could not find a unique line to replace")
                continue
            raw = raw.replace(old_line, new_line)
        json.loads(raw)
        if not a.dry_run:
            open(path, "w", encoding="utf-8", newline="").write(raw)

    files = {}
    for path, lang, key, old, new in app:
        full = os.path.join(ROOT, path)
        text = files.get(full) or open(full, encoding="utf-8", newline="").read()
        literals = [f"'{old}'", f'"{old}"']
        found = [l for l in literals if text.count(l) == 1]
        if not old or not found:
            problems.append(f"{path} {key} [{lang}]: current text not found exactly once; fix by hand")
            continue
        quote = found[0][0]
        replacement = quote + new.replace(quote, "\\" + quote) + quote
        print(f"app [{lang}] {key}: {old!r} -> {new!r}")
        files[full] = text.replace(found[0], replacement)
    if not a.dry_run:
        for full, text in files.items():
            open(full, "w", encoding="utf-8", newline="").write(text)

    if problems:
        print("\nNeeds a manual fix:")
        for p in problems:
            print("  -", p)
    print("\n(dry run, nothing written)" if a.dry_run else "\nDone. Run the web and app tests before committing.")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
